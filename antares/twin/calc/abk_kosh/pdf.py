# -*- coding: utf-8 -*-
"""PDF кошторису СЕС АБК (Deye SUN-100K-G03) із перерахованого xlsx — ті самі числа, що в Excel."""
import openpyxl, asyncio, html, json
from playwright.async_api import async_playwright
D = '/home/claude/work/abk_kosh/'
RJ = json.load(open(D + 'rows.json')); XN = RJ['XN']; KR = RJ['KR']; ROW = RJ['ROW']
wb = openpyxl.load_workbook(D + XN + '.xlsx', data_only=True)
K, P, T = wb['Кошторис'], wb['Параметри'], wb['Стрінги']
e = lambda x: html.escape(str(x)) if x is not None else ''
def nf(x, d=0):
    if x is None or x == '': return ''
    if isinstance(x, str): return e(x)
    return f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')
def qty(x):
    if isinstance(x, (int, float)) and abs(x - round(x)) > 1e-9: return nf(x, 2)
    return nf(x)
def pct(x): return (nf(x * 100) if abs(x * 100 - round(x * 100)) < 1e-9 else nf(x * 100, 1)) + ' %'
h = lambda k: K[f'H{KR[k]}'].value; i_ = lambda k: K[f'I{KR[k]}'].value
pa = lambda k: P[f'B{ROW[k]}'].value; pd = lambda k: P[f'C{ROW[k]}'].value
rows = []
last = KR['pkwp']
for r in range(5, last + 1):
    v = [K.cell(r, c).value for c in range(1, 11)]
    if all(a is None for a in v): continue
    no, nm, u, qA, qD, pr, mt, sA, sD, src = v
    if no and nm is None:                                                        # розділ
        rows.append(f'<tr class="sec"><td colspan="9">{e(no)}</td></tr>'); continue
    if no is None:
        cls = 'grand' if (nm.startswith('СОБІВАРТІСТЬ') or 'ПІД КЛЮЧ' in nm) else ('vat' if nm.startswith('ПДВ') else 'tot')
        d = 1 if nm.startswith('На 1 кВтп') else 0
        rows.append(f'<tr class="{cls}"><td></td><td colspan="5">{e(nm)}</td><td class="n">{nf(sA)}</td><td class="n">{nf(sD)}</td><td></td></tr>'); continue
    isp = qA is None and isinstance(pr, float) and pr < 1.5
    prs = (pct(pr) if pr < 1 else nf(pr, 4)) if isp else (nf(pr, 2) if isinstance(pr, float) and abs(pr - round(pr)) > 1e-9 and pr < 1000 else nf(pr))
    dstr = '' if no in ('6.1', '6.2', '2.1', '3.3', '3.4') or not isinstance(mt, (int, float)) or mt == 0 else pct(mt)
    dif = ' df' if isinstance(sA, (int, float)) and isinstance(sD, (int, float)) and abs(sA - sD) > 0.5 and no.startswith('4.') and (sA == 0) else ''
    rows.append(f'<tr class="ln{dif}"><td>{e(no)}</td><td>{e(nm)}</td><td class="c">{e(u)}</td><td class="n">{qty(qA) if qA is not None else ""}</td><td class="n">{qty(qD) if qD is not None else ""}</td>'
                f'<td class="n{" gap" if pr == 0 else ""}">{prs if pr != 0 else "—"}</td><td class="n b">{nf(sA) if sA else "—"}</td><td class="n b">{nf(sD) if sD else "—"}</td><td class="bs">{e(src)}{(" · мито " + dstr) if dstr else ""}</td></tr>')
SEC = [('Закупівля в Китаї', 't1'), ('Комісія', 't2'), ('Доставка, страхування, мито', 't3'), ('Купуємо в Україні', 't4'), ('Роботи', 't5'), ('Запас і винагорода', 't6')]
srow = ''.join(f'<tr><td>{a}</td><td class="n">{nf(h(k))}</td><td class="n">{nf(i_(k))}</td><td class="n d">{("+" if i_(k) - h(k) >= 0 else "−") + nf(abs(i_(k) - h(k)))}</td></tr>' for a, k in SEC)
eff = ''
for nm, k, d in (('Генерація модулів за рік (DC), МВт·год', 'e_dc', 1), ('На виході інвертора, МВт·год', 'e_ac', 1), ('Питома генерація, кВт·год/кВтп', 'e_y', 0),
                 ('Економія за перший рік, $ без ПДВ', 'e_rev', 0), ('Проста окупність, років', 'e_pb', 1), ('Під ключ на 1 МВт·год за рік, $', 'e_mwh', 0)):
    eff += f'<tr><td>{nm}</td><td class="n">{nf(h(k), d)}</td><td class="n">{nf(i_(k), d)}</td></tr>'
# стрінги
r2 = RJ['r2']; chk = ''
for j in (1, 2):
    v = [T.cell(r2 + j, c).value for c in range(1, 13)]
    ok = lambda s: f'<span class="{"ok" if s == "так" else "no"}">{"✓" if s == "так" else "✗"}</span>'
    chk += (f'<tr><td>{e(v[0])}</td><td class="n">{v[1]}</td><td class="n">{nf(v[2], 1)} {ok(v[3])}</td><td class="n">{nf(v[4], 1)} {ok(v[5])}</td>'
            f'<td class="n">{nf(v[6], 1)} {ok(v[7])}</td><td class="n">{nf(v[8], 1)} {ok(v[9])}</td><td class="n">{nf(v[10], 1)} {ok(v[11])}</td></tr>')
st = {'«Як є»': '', '«ДБН»': ''}
r = r2 + 5
while r <= T.max_row:
    v = [T.cell(r, c).value for c in range(1, 11)]
    if v[0] in st:
        st[v[0]] += (f'<tr><td>{e(v[1])}</td><td class="c">{v[2]}</td><td class="n">{v[3]}</td><td>{e(v[4])}</td><td>{e(v[5])}</td><td class="n">{nf(v[6], 1)}</td>'
                     f'<td class="n">{nf(v[7], 2)}</td><td class="n">{nf(v[8], 1)}</td><td class="n">{nf(v[9])}</td></tr>')
    elif isinstance(v[0], str) and v[0].startswith('Разом'):
        k = v[0].replace('Разом ', '')
        st[k] += f'<tr class="tot"><td>Разом</td><td></td><td class="n">{v[3]}</td><td></td><td></td><td></td><td class="n">{nf(v[7], 1)}</td><td></td><td class="n">{nf(v[9])}</td></tr>'
    r += 1
notes = []
r = last + 1
while r <= K.max_row:
    t = K[f'B{r}'].value
    if isinstance(t, str) and r > KR['e_mwh']: notes.append(t)
    r += 1
TH = '<tr><th>Стрінг</th><th class="c">MPPT</th><th class="n">Модулів</th><th>Де</th><th>Грані</th><th class="n">Нахил, °</th><th class="n">МВт·год</th><th class="n">Невідп., %</th><th class="n">DC, м</th></tr>'
HT = f'''<!doctype html><html lang="uk"><head><meta charset="utf-8"><title>Кошторис СЕС АБК Антарес</title><style>
@page {{ size:A4 landscape; margin:10mm 11mm 11mm 12mm; }}
body {{ font-family:'DejaVu Sans',sans-serif; font-size:7.2pt; line-height:1.25; color:#1d2a33; margin:0; }}
h1 {{ font-size:12.5pt; margin:0 0 2px; color:#123044; }} h2 {{ font-size:9.6pt; margin:9px 0 4px; color:#123044; }}
.sub {{ color:#5b6b76; font-size:7.3pt; margin-bottom:6px; }}
table {{ border-collapse:collapse; width:100%; }} td, th {{ border-bottom:0.4pt solid #d5dde2; padding:1.6px 4px; vertical-align:top; }}
th {{ background:#e8eef2; text-align:left; font-weight:bold; font-size:7pt; border-bottom:0.8pt solid #9fb1bd; }}
.n {{ text-align:right; white-space:nowrap; }} .c {{ text-align:center; }} .b {{ font-weight:bold; }} .d {{ color:#1f5f78; }}
.bs {{ color:#4d5c66; font-size:6.7pt; }}
tr.sec td {{ background:#f2f5f7; font-weight:bold; color:#123044; padding-top:4px; border-bottom:0.6pt solid #9fb1bd; }}
tr.tot td {{ font-weight:bold; border-bottom:0.8pt solid #9fb1bd; }}
tr.grand td {{ font-weight:bold; background:#e3edf3; font-size:8pt; border-bottom:0.8pt solid #6f8796; }}
tr.vat td {{ color:#4d5c66; }} tr.df td {{ background:#f6fbf8; }}
.kpi {{ display:flex; gap:8px; margin:6px 0 8px; }} .k {{ flex:1; border:0.6pt solid #c9d4db; border-radius:4px; padding:5px 7px; }}
.k b {{ display:block; font-size:10.5pt; color:#123044; }} .k span {{ color:#5b6b76; font-size:6.8pt; }} .k i {{ font-style:normal; color:#123044; font-weight:bold; }}
.two {{ display:flex; gap:14px; }} .two > div {{ flex:1; }}
.note {{ color:#4d5c66; font-size:6.9pt; margin:3px 0; }}
.ok {{ color:#1d6b40; font-weight:bold; }} .no {{ color:#b3261e; font-weight:bold; }}
.pb {{ page-break-before:always; }} td.gap {{ background:#fff3c4; }}
</style></head><body>
<h1>Кошторис будівництва під ключ — СЕС на даху АБК, ТОВ МК «Антарес»</h1>
<div class="sub">Інвертор Deye SUN-100K-G03 (100 кВт, 6 MPPT, DC до 1000 В), модулі ZNSHINE ZXN8-BD132-730. Два варіанти: «Як є» — розкладка К1 з доміком на надбудові; «ДБН» — з драбинами П1, огородженням 0,6 м і ходовими доріжками. 28.10.2026, $ без ПДВ. E-Systems.</div>
<div class="kpi">
<div class="k"><span>Модулів / потужність</span><b>{nf(pa("N"))} / {nf(pd("N"))} шт.</b><span>«Як є» {nf(pa("KWP"), 2)} кВтп · «ДБН» {nf(pd("KWP"), 2)} кВтп</span></div>
<div class="k"><span>Під ключ без ПДВ</span><b>{nf(h("tot"))} / {nf(i_("tot"))} $</b><span>«Як є» / «ДБН»</span></div>
<div class="k"><span>Під ключ з ПДВ 20 %</span><b>{nf(h("totv"))} / {nf(i_("totv"))} $</b><span>з пільгою «245»: {nf(h("tot245"))} / {nf(i_("tot245"))} $</span></div>
<div class="k"><span>На 1 кВтп без ПДВ</span><b>{nf(h("pkwp"))} / {nf(i_("pkwp"))} $</b><span>«Як є» / «ДБН»</span></div>
<div class="k"><span>Генерація / окупність</span><b>{nf(h("e_ac"), 1)} / {nf(i_("e_ac"), 1)} МВт·год</b><span>окупність {nf(h("e_pb"), 1)} / {nf(i_("e_pb"), 1)} року, без ПДВ</span></div>
</div>
<div class="two"><div>
<h2>Підсумок за розділами, $ без ПДВ</h2>
<table><tr><th>Розділ</th><th class="n">«Як є»</th><th class="n">«ДБН»</th><th class="n">Різниця</th></tr>{srow}
<tr class="grand"><td>Разом під ключ без ПДВ</td><td class="n">{nf(h("tot"))}</td><td class="n">{nf(i_("tot"))}</td><td class="n d">{("+" if i_("tot") - h("tot") >= 0 else "−") + nf(abs(i_("tot") - h("tot")))}</td></tr></table>
</div><div>
<h2>Ефект для АБК</h2>
<table><tr><th>Показник</th><th class="n">«Як є»</th><th class="n">«ДБН»</th></tr>{eff}</table>
<div class="note">1 кВт·год — {nf(pa("TAR"), 2)} грн з ПДВ ({nf(h("e_val") * 100, 2)} ¢ без ПДВ), {nf(pa("FX"), 2)} грн/$; ККД інвертора {pct(pa("ETA"))}. Генерація — повна модель K1: Meteonorm 8.1, тіні DSM, байпас-діоди, невідповідність модулів у стрінгу.</div>
</div></div>
<h2>Кошторис по позиціях</h2>
<table><tr><th style="width:3.5%">№</th><th style="width:36%">Позиція</th><th class="c" style="width:4%">Од.</th><th class="n" style="width:6%">«Як є»</th><th class="n" style="width:6%">«ДБН»</th><th class="n" style="width:6.5%">Ціна, $</th><th class="n" style="width:7.5%">Сума «Як є»</th><th class="n" style="width:7.5%">Сума «ДБН»</th><th>Підстава</th></tr>
{''.join(rows)}</table>
<div class="pb"></div>
<h2>Стрінги і MPPT інвертора Deye SUN-100K-G03</h2>
<div class="sub">Модуль: Voc 49,5 В, Vmp 41,7 В, Imp 17,51 А, Isc 18,51 А, Voc −0,25 %/K. Інвертор: DC до 1000 В, старт 250 В, MPPT 200–850 В, 40 А (Isc 60 А) на MPPT, 6 MPPT по 2 стрінги. Стрінги — однієї грані й нахилу; стрінг стін — на окремому MPPT.</div>
<table><tr><th>Варіант</th><th class="n">Модулів у стрінгу</th><th class="n">Voc при −30 °C ≤ 1000 В</th><th class="n">Vmp при 70 °C ≥ 250 В</th><th class="n">Vmp при −10 °C ≤ 850 В</th><th class="n">Imp на MPPT ≤ 40 А</th><th class="n">Isc на MPPT ≤ 60 А</th></tr>{chk}</table>
<div class="two" style="margin-top:6px"><div>
<h2>«Як є» — {nf(pa("ST"))} стрінгів по {nf(pa("LS"))}</h2>
<table>{TH}{st["«Як є»"]}</table>
</div><div>
<h2>«ДБН» — {nf(pd("ST"))} стрінгів по {nf(pd("LS"))}</h2>
<table>{TH}{st["«ДБН»"]}</table>
</div></div>
<h2>Примітки</h2>
{''.join(f'<div class="note">• {e(n)}</div>' for n in notes)}
</body></html>'''
open(D + 'abk_kosh.html', 'w', encoding='utf-8').write(HT)
async def run():
    async with async_playwright() as p:
        b = await p.chromium.launch(); pg = await b.new_page(); await pg.goto('file://' + D + 'abk_kosh.html'); await pg.wait_for_timeout(400)
        await pg.pdf(path=D + XN + '.pdf', print_background=True, prefer_css_page_size=True); await b.close()
asyncio.run(run())
print('ok', D + XN + '.pdf')
