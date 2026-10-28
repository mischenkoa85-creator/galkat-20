# -*- coding: utf-8 -*-
"""Кошторис будівництва під ключ СЕС АБК «Антарес», лише варіант «ДБН», українською й англійською (рішення власника 28.10, 17:52).
Інвертор Deye SUN-100K-G03 (DC до 1 000 В), 132 модулі ZNSHINE 730 Вт у 11 стрінгах по 12.
Кількості — antares/abk/abk_dbn.json (calc/v3/abk_dbn.py) і специфікація antares/abk/abk_bom.json (bom.py).
Ставки Китаю, доставки, мита, баласту, робіт, запасу й винагороди — кошторис під ключ 27.10; Україна — ціни магазинів 28.10 (prices.json).
Аркуші: Кошторис, Специфікація, Стрінги, Параметри. Формули — в Excel, синє — вихідні дані."""
import json, math, os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = '/home/claude/work/abk_kosh/'
XN = 'Кошторис_будівництва_під_ключ_СЕС_АБК_Антарес_ДБН_інвертор_Deye_100_кВт_28.10.2026'
A = '/home/claude/galkat-20/antares/abk/'
DB = json.load(open(A + 'abk_dbn.json')); BOM = json.load(open(A + 'abk_bom.json'))
PU = json.load(open(os.path.join(HERE, 'prices.json')))          # ціни України: {key: [грн з ПДВ, «Назва — джерело»]}
NB = BOM['n']
wb = openpyxl.Workbook()
AR = 'Arial'; BLUE = Font(name=AR, size=9, color='0000FF'); BLK = Font(name=AR, size=9); GRN = Font(name=AR, size=9, color='008000'); B = Font(name=AR, size=9, bold=True)
EN = Font(name=AR, size=8.5, color='4D5C66', italic=True); BB = Font(name=AR, size=10, bold=True); H1 = Font(name=AR, size=12, bold=True); SUB = Font(name=AR, size=8, color='5B6770')
HDR = PatternFill('solid', fgColor='E6EBF0'); SEC = PatternFill('solid', fgColor='F3F5F7'); GR = PatternFill('solid', fgColor='E3EDF3'); GRY = PatternFill('solid', fgColor='F7F7F7'); YEL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='C4CCD3'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top'); CEN = Alignment(wrap_text=True, vertical='center')
USD = '#,##0;(#,##0);-'; USD2 = '#,##0.00;(#,##0.00);-'; QTY = '#,##0.##;(#,##0.##);-'; PCT = '0.0%'
K27 = 'кошторис під ключ 27.10 / turnkey estimate 27.10'
# ======================= Параметри / Parameters
P = wb.active; P.title = 'Параметри'
P['A1'] = 'Параметри кошторису СЕС АБК, варіант «ДБН» / Estimate parameters, admin building PV, DBN option'; P['A1'].font = H1
P['A2'] = 'Синій — вихідні дані, чорний — формули. $ без ПДВ. / Blue — inputs, black — formulas. USD excl. VAT.'; P['A2'].font = SUB
ROW = {}
def hdr(ws, r, cols, letters):
    for c, h in zip(letters, cols): ws[f'{c}{r}'] = h; ws[f'{c}{r}'].font = B; ws[f'{c}{r}'].fill = HDR; ws[f'{c}{r}'].border = BOX; ws[f'{c}{r}'].alignment = CEN
hdr(P, 4, ('Параметр / Parameter', 'Значення / Value', 'Од. / Unit', 'Джерело / Source'), 'ABCD')
n_ = BOM['n']
Q = [  # key, name uk / en, value, unit, src, fmt, kind
 ('NA', 'Модулів на арці / Modules on the arch roof', n_['NA'], 'шт', '3D-модель АБК, варіант «ДБН» / 3D model', QTY, 'in'),
 ('NL', 'Модулів на нижньому даху, «домік» 15° / Modules on the lower roof, dual-tilt 15°', n_['NL'], 'шт', f"33 місця, {len(DB['drop'])} поза стрінгами / 33 places, {len(DB['drop'])} outside strings", QTY, 'in'),
 ('NN', 'Модулів на надбудові сходів, «домік» 15° / Modules on the stair house, dual-tilt 15°', n_['NN'], 'шт', '3D-модель АБК / 3D model', QTY, 'in'),
 ('NW', 'Модулів на південних стінах / Modules on south walls', n_['NW'], 'шт', '3D-модель АБК / 3D model', QTY, 'in'),
 ('N', 'Модулів у стрінгах / Modules in strings', '=B5+B6+B7+B8', 'шт', '', QTY, 'f'),
 ('NS', 'Модулів до поставки (+2 %) / Modules to supply (+2 %)', '=ROUNDUP(B9*{SPARE},0)', 'шт', K27, QTY, 'f'),
 ('KWP', 'Потужність СЕС, DC / PV capacity, DC', '=B9*{PW}/1000', 'кВтп / kWp', '730 Вт × модулів / 730 W × modules', '#,##0.00', 'f'),
 ('ST', 'Стрінгів / Strings', len(DB['strings']), 'шт', 'по 2 на MPPT, 6 MPPT / 2 per MPPT, 6 MPPT', QTY, 'in'),
 ('LS', 'Модулів у стрінгу / Modules per string', DB['L'], 'шт', 'Voc при −30 °C ≤ 1 000 В / Voc at −30 °C ≤ 1000 V', QTY, 'in'),
 ('MWH', 'Генерація модулів за рік (DC) / Annual module yield (DC)', DB['mwh'], 'МВт·год / MWh', 'повна модель K1: Meteonorm 8.1, тіні DSM, байпас-діоди / full K1 model', '#,##0.0', 'in'),
 ('TRB', 'Блоків під опори лотків / Blocks under tray supports', BOM['blk']['tray'], 'шт', f"лотки на даху {DB['tray_stats']['roof']} м, опора через 1,5 м / roof trays, support every 1.5 m", QTY, 'in'),
 ('BLKN', 'Баластних блоків до замовлення (+1 %) / Ballast blocks to order (+1 %)', '=ROUNDUP((B5*{BLKA}+(B6+B7)*{BLKD}+B15)*1.01,0)', 'шт', 'арка — як дах цеху; «домік» — 4 на модуль; опори лотків / arch as shop roof; dual-tilt 4 per module; tray supports', QTY, 'f'),
 ('ACM', 'Кабель AC інвертор → щит СЕС / AC cable inverter → PV board', 45, 'м', 'по стіні і стояку сходів, + 10 % / along wall and stair riser, + 10 %', QTY, 'in'),
 ('RSM', 'Кабель RS485 лічильник → інвертор / RS485 cable meter → inverter', 55, 'м', 'уздовж кабелю AC / along AC cable', QTY, 'in'),
 ('LP1', 'Драбина П1 на арку: 4,8 + 1,0 м / Fixed ladder P1 to arch: 4.8 + 1.0 m', 5.8, 'м', 'ДБН В.1.1-7:2016 п. 8.8–8.10', QTY, 'in'),
 ('PLAT', 'Площадка з огорожею 0,6 м на арці / Landing with 0.6 m guard on arch', 1, 'шт', 'ДБН В.1.1-7:2016 табл. 5', QTY, 'in'),
 ('LPN', 'Драбини П1 на надбудову / Fixed ladders P1 to stair house', 5.6, 'м', '2,25 + 1,0 і 1,35 + 1,0 м; ДБН В.1.1-7:2016 п. 8.8', QTY, 'in'),
 ('RAIL', 'Огородження покрівлі 0,6 м / Roof guardrail 0.6 m', 99, 'м', 'ДБН В.1.1-7:2016 п. 8.12; ДБН В.2.6-220:2017 п. 6.1.2', QTY, 'in'),
 ('WALK', 'Ходові доріжки 0,8 м / Walkways 0.8 m', 63, 'м²', 'ДБН В.2.6-220:2017 п. 7.1.3', QTY, 'in'),
 ('HRS', 'Монтаж драбин, площадки, огородження і доріжок / Installation of ladders, landing, guardrail, walkways', 112, 'люд.·год / man-h', 'драбини 2 × 16, огородження 64, доріжки 16 / ladders 2 × 16, guardrail 64, walkways 16', QTY, 'in'),
]
R0 = 5 + len(Q) + 1
SKEYS = ['PW', 'FX', 'SPARE', 'PMOD', 'MFL', 'MDOM', 'MWAL', 'DCS', 'TRS', 'BZM', 'RS', 'COM', 'DMOD', 'DOTH', 'INS', 'DUTK', 'BLK', 'BLKA', 'BLKD']
SR0 = {k: f'$B${R0 + 1 + i}' for i, k in enumerate(SKEYS)}
fq = lambda v: v.format(**SR0) if isinstance(v, str) and '{' in v else v
for i, (k, nm, v, u, src, fmt, kind) in enumerate(Q):
    r = 5 + i; ROW[k] = r
    P[f'A{r}'] = nm; P[f'B{r}'] = fq(v); P[f'C{r}'] = u; P[f'D{r}'] = src
    P[f'A{r}'].font = BLK; P[f'B{r}'].font = BLUE if kind == 'in' else BLK; P[f'B{r}'].number_format = fmt; P[f'C{r}'].font = BLK; P[f'D{r}'].font = SUB
    for c in 'ABCD': P[f'{c}{r}'].border = BOX
assert ROW['N'] == 9 and ROW['TRB'] == 15 and ROW['BLKN'] == 16
hdr(P, R0, ('Ставки і ціни / Rates and prices', 'Значення / Value', 'Од. / Unit', 'Джерело / Source'), 'ABCD')
S = [
 ('PW', 'Потужність модуля / Module power', 730, 'Вт / W', 'паспорт ZNSHINE / datasheet', QTY, 'in'),
 ('FX', 'Курс / Exchange rate', 44.89, 'грн/$', 'як у кошторисі 27.10 / as estimate 27.10', '#,##0.00', 'in'),
 ('SPARE', 'Запас модулів / Module spare', 1.02, '', K27, '0.00', 'in'),
 ('PMOD', 'Фотомодуль ZNSHINE ZXN8-BD132-730, FOB / PV module, FOB', 73, '$/шт', 'КП ES-261026-01', USD2, 'in'),
 ('MFL', 'Кріплення по даху на баласті, на модуль / Roof rail mounting on ballast, per module', '=21500/1925', '$/модуль', 'КП 26.10: 21 500 $ на 1 925 модулів / quote 26.10', USD2, 'f'),
 ('MDOM', 'Кріплення «домік» 15°, на модуль / Dual-tilt 15° mounting, per module', '=11211/1260+4*0.5', '$/модуль', 'ГалКат 15° 22.10 + 4 притиски по 0,50 $ / GalKat 15° + 4 clamps', USD2, 'f'),
 ('MWAL', 'Кріплення на стіні, на модуль / Wall mounting, per module', '=2555/297', '$/модуль', f'{K27}, поз. 1.3.2', USD2, 'f'),
 ('DCS', 'Кабелі DC, MC4, металорукав, на стрінг / DC cables, MC4, conduit, per string', '=11520/192', '$/стрінг', f'{K27}, поз. 1.4', USD2, 'f'),
 ('TRS', 'Лотки, фасонні, кріплення, на стрінг / Trays, fittings, supports, per string', '=12398/192', '$/стрінг', f'{K27}, поз. 1.7', USD2, 'f'),
 ('BZM', 'Блискавкозахист: ПЗІП, провідники, затискачі, на модуль / Lightning protection, per module', '=2400/2112', '$/модуль', f'{K27}, поз. 1.8', USD2, 'f'),
 ('RS', 'Кабель RS485 2×2×0,75 в екрані, FOB / RS485 cable, FOB', '=276/1050', '$/м', 'специфікація Китай 26.10 / China spec 26.10', '0.0000', 'f'),
 ('COM', 'Наша комісія від закупівлі в Китаї / Our commission on China purchase', 0.05, '', 'договір комісії / commission agreement', PCT, 'in'),
 ('DMOD', "Доставка модулів, 40'HQ / 594 модулі / Module freight per module", '=7000/594', '$/шт', f'{K27}, поз. 3.2', USD2, 'f'),
 ('DOTH', 'Доставка решти вантажу, на модуль / Other freight per module', '=14000/2112', '$/модуль', f'{K27}, поз. 3.3', USD2, 'f'),
 ('INS', 'Страхування вантажу, крім модулів / Cargo insurance except modules', 0.003, '', K27, PCT, 'in'),
 ('DUTK', 'Коефіцієнт митної вартості / Customs value factor', '=4840.7596950824/4234.64192', '', f'{K27}: мито / ставки × FOB', '0.0000', 'f'),
 ('BLK', 'Баластний блок з муфтою і матом / Ballast block with coupling and mat', 3, '$/шт', 'рішення власника 26.10 / owner decision 26.10', USD2, 'in'),
 ('BLKA', 'Блоків на модуль арки / Blocks per arch module', '=(4457+795)/1639', 'шт', f'{K27}, поз. 4.1–4.2', '0.000', 'f'),
 ('BLKD', 'Блоків на модуль «доміку» / Blocks per dual-tilt module', 4, 'шт', '4 опори на модуль / 4 supports per module', '0.000', 'in'),
 ('BLKDEL', "Доставка матеріалів України, на блок / Ukrainian materials delivery, per block", '=2107/6227', '$/шт', f'{K27}, поз. 4.9', '0.0000', 'f'),
 ('WRK', 'Встановлення і наладка СЕС / PV installation and commissioning', 50, '$/кВтп', 'рішення власника 26.10 / owner decision 26.10', USD2, 'in'),
 ('MECH', 'Механізація, на модуль / Cranes and lifts, per module', '=3080/2112', '$/модуль', f'{K27}, поз. 5.2', USD2, 'f'),
 ('CONS', 'Витратні матеріали, на модуль / Consumables, per module', '=2584/2112', '$/модуль', f'{K27}, поз. 5.8', USD2, 'f'),
 ('HRATE', 'Ставка монтажника / Installer rate', 700, 'грн/люд.·год', 'ставка власника для КП / owner rate', '#,##0', 'in'),
 ('RES', 'Запас / Contingency', 0.04, '', 'рішення власника 26.10', PCT, 'in'),
 ('GEN', 'Винагорода генпідрядника / General contractor fee', 0.18, '', 'рішення власника 26.10', PCT, 'in'),
 ('VAT', 'ПДВ / VAT', 0.2, '', 'ПКУ, ст. 193 / Tax Code art. 193', PCT, 'in'),
 ('TAR', 'Ціна 1 кВт·год, з ПДВ / Electricity price incl. VAT', 7.00, 'грн/кВт·год', 'оцінка власника / owner estimate', '#,##0.00', 'in'),
 ('ETA', 'ККД інвертора (європейський) / Inverter Euro efficiency', PU['eta'][0], '', PU['eta'][1], PCT, 'in'),
]
for k, (uah, src) in PU.items():
    if k == 'eta': continue
    nm, s_ = (src.split(' — ', 1) + [''])[:2]
    S.append(('U_' + k, nm, uah, 'грн з ПДВ / UAH incl. VAT', s_, '#,##0.00', 'in'))
for i, (k, nm, v, u, src, fmt, kind) in enumerate(S):
    r = R0 + 1 + i; ROW[k] = r
    P[f'A{r}'] = nm; P[f'B{r}'] = v; P[f'C{r}'] = u; P[f'D{r}'] = src
    P[f'A{r}'].font = BLK; P[f'B{r}'].font = BLUE if kind == 'in' else BLK; P[f'B{r}'].number_format = fmt; P[f'C{r}'].font = BLK; P[f'D{r}'].font = SUB
    for c in 'ABCD': P[f'{c}{r}'].border = BOX
assert all(SR0[k] == f'$B${ROW[k]}' for k in SKEYS)
for c, w in zip('ABCD', (70, 14, 18, 80)): P.column_dimensions[c].width = w
for r in range(5, R0 + 1 + len(S)): P[f'A{r}'].alignment = WR; P[f'D{r}'].alignment = WR
P.freeze_panes = 'A5'
pa = lambda k: f"Параметри!$B${ROW[k]}"
usd = lambda k: f"Параметри!$B${ROW['U_' + k]}/(1+{pa('VAT')})/{pa('FX')}"
src_of = lambda k: PU[k][1].split(' — ', 1)[-1]
# ======================= Кошторис / Estimate
K = wb.create_sheet('Кошторис', 0)
K['A1'] = 'Кошторис будівництва під ключ — СЕС на даху АБК, ТОВ МК «Антарес», варіант «ДБН»'; K['A1'].font = H1
K['A2'] = 'Turnkey construction estimate — rooftop PV plant on the administrative building, Antares LLC, DBN-compliant option'; K['A2'].font = Font(name=AR, size=11, bold=True, color='4D5C66')
K['A3'] = (f"Інвертор Deye SUN-100K-G03 (100 кВт, DC до 1 000 В), {NB['N']} модулів ZNSHINE 730 Вт у {NB['ST']} стрінгах по {DB['L']}. $ без ПДВ. "
           f"/ Inverter Deye SUN-100K-G03 (100 kW, DC up to 1000 V), {NB['N']} ZNSHINE 730 W modules in {NB['ST']} strings of {DB['L']}. USD excl. VAT."); K['A3'].font = SUB
H = ('№', 'Позиція', 'Item', 'Од. / Unit', 'К-сть / Qty', 'Ціна, $ / Price, $', 'Мито / Duty', 'Сума, $ / Amount, $', 'Підстава / Basis')
hdr(K, 5, H, 'ABCDEFGHI'); K.row_dimensions[5].height = 28
r = 6; KR = {}
def sec(uk, en):
    global r
    K[f'A{r}'] = uk; K[f'C{r}'] = en; K[f'A{r}'].font = B; K[f'C{r}'].font = Font(name=AR, size=9, bold=True, color='4D5C66')
    for c in 'ABCDEFGHI': K[f'{c}{r}'].fill = SEC; K[f'{c}{r}'].border = BOX
    r += 1
def line(no, uk, en, u, q, pr, duty, src, key=None, s=None, pfmt=USD2):
    global r
    K[f'A{r}'] = no; K[f'B{r}'] = uk; K[f'C{r}'] = en; K[f'D{r}'] = u; K[f'E{r}'] = q; K[f'F{r}'] = pr; K[f'G{r}'] = duty; K[f'I{r}'] = src
    K[f'H{r}'] = s if s else f'=E{r}*F{r}'
    K[f'A{r}'].font = BLK; K[f'B{r}'].font = BLK; K[f'C{r}'].font = EN; K[f'D{r}'].font = BLK; K[f'I{r}'].font = SUB
    K[f'E{r}'].font = GRN if isinstance(q, str) and 'Параметри' in q else (BLUE if isinstance(q, (int, float)) else BLK); K[f'E{r}'].number_format = QTY
    K[f'F{r}'].font = GRN if isinstance(pr, str) and 'Параметри' in pr else (BLUE if isinstance(pr, (int, float)) else BLK); K[f'F{r}'].number_format = pfmt
    K[f'G{r}'].font = BLUE; K[f'G{r}'].number_format = '0.0%;;-'; K[f'H{r}'].font = BLK; K[f'H{r}'].number_format = USD
    for c in 'ABCDEFGHI': K[f'{c}{r}'].border = BOX
    for c in 'BCI': K[f'{c}{r}'].alignment = WR
    if key: KR[key] = r
    r += 1
def total(uk, en, f, key, grand=False):
    global r
    K[f'B{r}'] = uk; K[f'C{r}'] = en; K[f'H{r}'] = f
    for c in 'ABCDEFGHI':
        K[f'{c}{r}'].border = BOX
        if grand: K[f'{c}{r}'].fill = GR
    K[f'B{r}'].font = BB if grand else B; K[f'C{r}'].font = Font(name=AR, size=9, bold=True, color='4D5C66'); K[f'H{r}'].font = BB if grand else B; K[f'H{r}'].number_format = USD
    K[f'B{r}'].alignment = WR; K[f'C{r}'].alignment = WR
    KR[key] = r; r += 1
q = lambda k: f'=Параметри!B{ROW[k]}'
sec('1. Закупівля в Китаї, FOB — разом із замовленням СЕС цеху', '1. Purchase in China, FOB — together with the shop roof PV order')
s1 = r
line('1.1', 'Фотомодулі ZNSHINE ZXN8-BD132-730 (730 Вт, виводи 1,2–1,4 м), з запасом 2 %', 'PV modules ZNSHINE ZXN8-BD132-730 (730 W, leads 1.2–1.4 m), 2 % spare', 'шт / pcs', q('NS'), f'={pa("PMOD")}', 0, f'{K27}, поз. 1.2', 'mod')
line('1.2', 'Кріплення на арці по даху: лаги, гільзи, шпильки, профілі, кутники, притиски, метизи', 'Arch roof mounting: rails, sleeves, rods, profiles, angles, clamps, fasteners', 'модуль / module', q('NA'), f'={pa("MFL")}', 0.008, 'КП 26.10, на модуль / per module', 'mfl')
line('1.3', 'Кріплення «домік» 15° на нижньому даху і надбудові: трикутники, кронштейни, притиски, метизи', 'Dual-tilt 15° mounting on lower roof and stair house: triangles, brackets, clamps, fasteners', 'модуль / module', f'={q("NL")[1:]}+{q("NN")[1:]}', f'={pa("MDOM")}', 0.008, 'ГалКат 15° 22.10, на модуль / per module', 'mdom')
line('1.4', 'Кріплення на південних стінах: рейки, кронштейни з анкерами, притиски', 'South wall mounting: rails, brackets with anchors, clamps', 'модуль / module', q('NW'), f'={pa("MWAL")}', 0.008, f'{K27}, поз. 1.3.2', 'mwal')
line('1.5', "Кабелі DC H1Z2Z2-K 1×6 мм², роз'єми MC4, металорукав, стяжки, бирки", 'DC cables H1Z2Z2-K 1×6 mm², MC4 connectors, conduit, ties, labels', 'стрінг / string', q('ST'), f'={pa("DCS")}', 0.067, f'{K27}, поз. 1.4', 'dc')
line('1.6', 'Лотки 100×50 з кришками, фасонні, опори, кронштейни, з\'єднувачі', 'Trays 100×50 with lids, fittings, supports, brackets, couplers', 'стрінг / string', q('ST'), f'={pa("TRS")}', 0.001, f'{K27}, поз. 1.7', 'tray')
line('1.7', 'Блискавкозахист і вирівнювання потенціалів: ПЗІП RS485, провідники 1×6 і 1×16, затискачі, контактні шайби', 'Lightning protection and bonding: RS485 SPD, 1×6 and 1×16 conductors, clamps, contact washers', 'модуль / module', q('N'), f'={pa("BZM")}', 0.05, f'{K27}, поз. 1.8', 'bz')
line('1.8', 'Кабель RS485 2×2×0,75 в екрані — лічильник → інвертор', 'RS485 cable 2×2×0.75 screened — meter to inverter', 'м / m', q('RSM'), f'={pa("RS")}', 0.10, 'специфікація Китай 26.10 / China spec', 'rs', pfmt='0.00')
e1 = r - 1
total('Разом закупівля в Китаї', 'Total purchase in China', f'=SUM(H{s1}:H{e1})', 't1')
sec('2. Наша комісія — супровід замовлення за договором комісії', '2. Our commission — order support under commission agreement')
line('2.1', 'Комісія від закупівлі в Китаї', 'Commission on China purchase', '', None, f'={pa("COM")}', None, 'договір комісії / agreement', 'com', f'=H{KR["t1"]}*F{r}', PCT)
total('Разом комісія', 'Total commission', f'=H{KR["com"]}', 't2')
sec("3. Доставка до об'єкта, страхування, мито", '3. Delivery to site, insurance, customs duty')
s3 = r
line('3.1', "Доставка модулів у контейнерах 40'HQ", "Module freight in 40'HQ containers", 'шт / pcs', q('NS'), f'={pa("DMOD")}', None, f'{K27}, поз. 3.2', 'dmod')
line('3.2', 'Доставка кріплення, кабелів і лотків', 'Freight of mounting, cables and trays', 'модуль / module', q('N'), f'={pa("DOTH")}', None, f'{K27}, поз. 3.3', 'doth')
line('3.3', 'Страхування вантажу, крім модулів', 'Cargo insurance, except modules', '', None, f'={pa("INS")}', None, K27, 'ins', f'=(H{KR["t1"]}-H{KR["mod"]})*F{r}', '0.0%')
line('3.4', 'Мито на товари, крім модулів (модулі — пільга «245»)', 'Customs duty except modules (modules — "245" exemption)', '', None, f'={pa("DUTK")}', None, 'ставки за кодами УКТ ЗЕД / tariff rates', 'duty', f'=SUMPRODUCT(G{s1}:G{e1},H{s1}:H{e1})*F{r}', '0.0000')
e3 = r - 1
total('Разом доставка і мито', 'Total delivery and duty', f'=SUM(H{s3}:H{e3})', 't3')
sec('4. Купуємо в Україні: інвертор, підключення, баласт, безпека покрівлі', '4. Purchased in Ukraine: inverter, connection, ballast, roof safety')
s4 = r
line('4.1', "Інвертор мережевий Deye SUN-100K-G03: 100 кВт, 3 фази 400 В, 6 MPPT, DC до 1 000 В, IP65; DC-роз'єднувач, ПЗІП тип II на DC і AC", 'Grid inverter Deye SUN-100K-G03: 100 kW, 3-phase 400 V, 6 MPPT, DC up to 1000 V, IP65; DC switch, type II SPD on DC and AC', 'шт / pcs', 1, f'={usd("deye")}', None, src_of('deye'), 'inv')
line('4.2', 'Баластний блок 300×300×60 з муфтою і матом — арка, «домік», опори лотків (+1 %)', 'Ballast block 300×300×60 with coupling and mat — arch, dual-tilt, tray supports (+1 %)', 'шт / pcs', q('BLKN'), f'={pa("BLK")}', None, 'рішення власника 26.10 / owner decision', 'blk')
line('4.3', 'Кабель ВВГнг 5×70 мм², мідь — інвертор → щит СЕС біля вступного щита АБК', 'Cable VVGng 5×70 mm² copper — inverter to PV board at the building main board', 'м / m', q('ACM'), f'={usd("ac")}', None, src_of('ac'), 'ac')
line('4.4', 'Щит СЕС: автоматичний вимикач ABB A2B 250 TMF 200 (3P, 200 А) у корпусі IP54', 'PV board: MCCB ABB A2B 250 TMF 200 (3P, 200 A) in IP54 enclosure', 'компл / set', 1, f'={usd("mccb")}+{usd("box")}', None, src_of('mccb') + '; ' + src_of('box'), 'mccb')
line('4.5', 'ПЗІП AC тип 2, 4P (Eaton SPCT2-280/4) у щиті СЕС', 'AC surge protector type 2, 4P (Eaton SPCT2-280/4) in PV board', 'шт / pcs', 1, f'={usd("spd")}', None, src_of('spd'), 'spd')
line('4.6', 'Лічильник CHINT DTSU666 з трансформаторами струму 250 А — обмеження видачі в мережу (Zero Export)', 'Meter CHINT DTSU666 with 250 A CTs — export limiting (Zero Export)', 'компл / set', 1, f'={usd("meter")}', None, src_of('meter'), 'meter')
line('4.7', "Доставка матеріалів України на об'єкт", 'Delivery of Ukrainian materials to site', 'блок / block', q('BLKN'), f'={pa("BLKDEL")}', None, f'{K27}, поз. 4.9', 'del')
line('4.8', 'Драбина П1 на арку з поручнями над парапетом', 'Fixed ladder P1 to the arch with handrails above parapet', 'м / m', q('LP1'), f'={usd("lp1")}', None, src_of('lp1'), 'lp1')
line('4.9', 'Площадка з огорожею 0,6 м на арці перед виходом на покрівлю', 'Landing with 0.6 m guard on the arch at the roof exit', 'шт / pcs', q('PLAT'), 0, None, 'ціна — за КП виробника драбини / price per ladder maker quote', 'plat', pfmt=USD)
K[f'F{KR["plat"]}'].fill = YEL
line('4.10', 'Драбини П1 на надбудову сходів (2 шт.)', 'Fixed ladders P1 to the stair house (2 pcs)', 'м / m', q('LPN'), f'={usd("lp1")}', None, src_of('lp1'), 'lpn')
line('4.11', 'Огородження покрівлі 0,6 м, металеве', 'Roof guardrail 0.6 m, steel', 'м / m', q('RAIL'), f'={usd("rail")}', None, src_of('rail'), 'rail')
line('4.12', 'Ходові доріжки 0,8 м — гумові плити 500×500×20 мм', 'Walkways 0.8 m — rubber tiles 500×500×20 mm', 'м² / m²', q('WALK'), f'={usd("walk")}', None, src_of('walk'), 'walk')
e4 = r - 1
total('Разом купуємо в Україні', 'Total purchased in Ukraine', f'=SUM(H{s4}:H{e4})', 't4')
sec('5. Роботи', '5. Works')
s5 = r
line('5.1', 'Встановлення і наладка СЕС, електромонтажні роботи DC і AC, налаштування Zero Export, вимірювання', 'PV installation and commissioning, DC and AC electrical works, Zero Export set-up, measurements', 'кВтп / kWp', q('KWP'), f'={pa("WRK")}', None, 'рішення власника 26.10 / owner decision', 'wrk')
line('5.2', 'Механізація: автокран і автовишка', 'Crane and aerial lift', 'модуль / module', q('N'), f'={pa("MECH")}', None, f'{K27}, поз. 5.2', 'mech')
line('5.3', 'Витратні матеріали: наконечники, маркування, прохід крізь стіну, ЗІЗ, вивезення упаковки', 'Consumables: lugs, marking, wall pass-through, PPE, packaging removal', 'модуль / module', q('N'), f'={pa("CONS")}', None, f'{K27}, поз. 5.8', 'cons')
line('5.4', 'Монтаж драбин П1, площадки, огородження з анкеруванням стійок, укладання доріжок', 'Installation of P1 ladders, landing, guardrail with anchored posts, walkways', 'люд.·год / man-h', q('HRS'), f'={pa("HRATE")}/{pa("FX")}', None, '700 грн/люд.·год / UAH per man-hour', 'hrs')
e5 = r - 1
total('Разом роботи', 'Total works', f'=SUM(H{s5}:H{e5})', 't5')
r += 1
total('СОБІВАРТІСТЬ (розділи 1–5)', 'COST (sections 1–5)', '=' + '+'.join(f'H{KR[t]}' for t in ('t1', 't2', 't3', 't4', 't5')), 'cost', True)
r += 1
sec('6. Запас і винагорода генпідрядника', '6. Contingency and general contractor fee')
line('6.1', 'Запас', 'Contingency', '', None, f'={pa("RES")}', None, 'рішення власника 26.10', 'res', f'=H{KR["cost"]}*F{r}', PCT)
line('6.2', 'Винагорода генпідрядника', 'General contractor fee', '', None, f'={pa("GEN")}', None, 'рішення власника 26.10', 'gen', f'=H{KR["cost"]}*F{r}', PCT)
total('Разом запас і винагорода', 'Total contingency and fee', f'=H{KR["res"]}+H{KR["gen"]}', 't6')
r += 1
total('РАЗОМ ПІД КЛЮЧ, $ без ПДВ', 'TOTAL TURNKEY, USD excl. VAT', f'=H{KR["cost"]}+H{KR["t6"]}', 'tot', True)
total('ПДВ 20 % на всю суму', 'VAT 20 % on the full amount', f'=H{KR["tot"]}*{pa("VAT")}', 'vat')
total('РАЗОМ ПІД КЛЮЧ, $ з ПДВ', 'TOTAL TURNKEY, USD incl. VAT', f'=H{KR["tot"]}+H{KR["vat"]}', 'totv', True)
r += 1
total('ПДВ 20 %, якщо модулі ввозяться на замовника за пільгою «245»', 'VAT 20 % if modules are imported by the customer under the "245" exemption', f'=(H{KR["tot"]}-H{KR["mod"]})*{pa("VAT")}', 'vat245')
total('РАЗОМ ПІД КЛЮЧ, $ з ПДВ (з пільгою «245»)', 'TOTAL TURNKEY, USD incl. VAT ("245" exemption)', f'=H{KR["tot"]}+H{KR["vat245"]}', 'tot245', True)
r += 1
total('На 1 кВтп СЕС, $ без ПДВ', 'Per 1 kWp, USD excl. VAT', f'=H{KR["tot"]}/Параметри!B{ROW["KWP"]}', 'pkwp')
r += 1
sec('7. Ефект СЕС для АБК', '7. Effect of the PV plant for the building')
def eff(uk, en, f, fmt, key, u=''):
    global r
    K[f'B{r}'] = uk; K[f'C{r}'] = en; K[f'D{r}'] = u; K[f'H{r}'] = f
    for c in 'ABCDEFGHI': K[f'{c}{r}'].border = BOX
    K[f'B{r}'].font = BLK; K[f'C{r}'].font = EN; K[f'D{r}'].font = BLK; K[f'H{r}'].font = BLK; K[f'H{r}'].number_format = fmt
    KR[key] = r; r += 1
eff('Потужність СЕС DC / AC інвертора 100 кВт', 'PV capacity DC / inverter AC 100 kW', f'=Параметри!B{ROW["KWP"]}', '#,##0.00', 'e_kwp', 'кВтп / kWp')
eff('Генерація модулів за рік (DC, повна модель K1)', 'Annual module yield (DC, full K1 model)', f'=Параметри!B{ROW["MWH"]}', '#,##0.0', 'e_dc', 'МВт·год / MWh')
eff('Генерація на виході інвертора за рік', 'Annual inverter AC output', f'=H{r-1}*{pa("ETA")}', '#,##0.0', 'e_ac', 'МВт·год / MWh')
eff('Питома генерація', 'Specific yield', f'=H{r-1}*1000/H{r-3}', '#,##0', 'e_y', 'кВт·год/кВтп')
eff('Ціна 1 кВт·год без ПДВ', 'Electricity price excl. VAT', f'={pa("TAR")}/(1+{pa("VAT")})/{pa("FX")}', '0.0000', 'e_val', '$/кВт·год')
eff('Економія за перший рік, $ без ПДВ', 'First-year savings, USD excl. VAT', f'=H{KR["e_ac"]}*1000*H{KR["e_val"]}', USD, 'e_rev', '$')
eff('Проста окупність, без ПДВ', 'Simple payback, excl. VAT', f'=H{KR["tot"]}/H{KR["e_rev"]}', '0.0', 'e_pb', 'років / years')
eff('Вартість під ключ на 1 МВт·год річної генерації', 'Turnkey cost per 1 MWh of annual yield', f'=H{KR["tot"]}/H{KR["e_ac"]}', USD, 'e_mwh', '$')
r += 1
NOTES = [
 (f"Стрінги: {NB['ST']} × {DB['L']} модулів, по 2 на MPPT, стрінг стін — на окремому MPPT; Voc при −30 °C — {DB['strings'][0]['voc30']:.0f} В ≤ 1 000 В. Повна специфікація — аркуш «Специфікація», стрінги — аркуш «Стрінги».",
  f"Strings: {NB['ST']} × {DB['L']} modules, 2 per MPPT, the wall string on its own MPPT; Voc at −30 °C {DB['strings'][0]['voc30']:.0f} V ≤ 1000 V. Full bill of materials — sheet «Специфікація», strings — sheet «Стрінги»."),
 ('Інвертор — на північній стіні надбудови сходів над нижнім дахом (+10,6), у тіні; кабель AC — у надбудову і стояком сходів до щита СЕС біля вступного щита АБК. Струм інвертора до 159,4 А при 400 В, вимикач 200 А, кабель 5×70 мм² мідь.',
  'Inverter on the north wall of the stair house above the lower roof (+10.6), in shade; AC cable into the stair house and down the stair riser to the PV board next to the building main board. Inverter current up to 159.4 A at 400 V, 200 A breaker, 5×70 mm² copper cable.'),
 ('Блискавкозахист — без зовнішніх щогл (як том 7 СЕС цеху): вирівнювання потенціалів кріплення провідником Cu 1×6 у лотках, по 2 приєднання на блок кріплення, Cu 1×16 до PE інвертора і до прутка на парапеті в місці переходу лотка; ПЗІП тип II в інверторі, тип II у щиті СЕС, ПЗІП RS485.',
  'Lightning protection without external masts (as volume 7 of the shop PV): bonding of the mounting with Cu 1×6 in trays, 2 connections per mounting block, Cu 1×16 to inverter PE and to the parapet rod where the tray crosses; type II SPD in the inverter and in the PV board, RS485 SPD.'),
 ('Робота без видачі в мережу: лічильник на вводі АБК і обмеження потужності інвертора (Zero Export); про генеруючу установку без видачі в мережу повідомляють ОСР за Додатком 11 до Кодексу систем розподілу.',
  'Operation without export: meter at the building input and inverter power limiting (Zero Export); the DSO is notified of a non-exporting generating unit per Annex 11 of the Distribution System Code.'),
 ('Не входить: обстеження будівлі й несучої здатності покрівлі (постанова КМУ № 406), проєктна документація, ремонт покрівлі; площадка з огорожею на арці (поз. 4.9) — за КП виробника драбини.',
  'Not included: building and roof load survey (CMU Resolution No. 406), design documentation, roof repair; the landing on the arch (item 4.9) — per ladder maker quote.'),
]
for uk, en in NOTES:
    K[f'B{r}'] = uk; K[f'B{r}'].font = SUB; K.merge_cells(f'B{r}:I{r}'); K[f'B{r}'].alignment = WR; K.row_dimensions[r].height = 26; r += 1
    K[f'B{r}'] = en; K[f'B{r}'].font = Font(name=AR, size=8, color='7A8790', italic=True); K.merge_cells(f'B{r}:I{r}'); K[f'B{r}'].alignment = WR; K.row_dimensions[r].height = 26; r += 1
for c, w in zip('ABCDEFGHI', (5.5, 52, 46, 11, 9, 11, 7, 13, 44)): K.column_dimensions[c].width = w
K.freeze_panes = 'D6'
# ======================= Специфікація / Bill of materials
SP = wb.create_sheet('Специфікація', 1)
SP['A1'] = 'Специфікація СЕС АБК, варіант «ДБН» / Bill of materials, admin building PV, DBN option'; SP['A1'].font = H1
SP['A2'] = 'Кількості — з 3D-моделі АБК і розрахунку стрінгів; ціни — у кошторисі за рядком. / Quantities from the 3D model and string design; prices in the estimate line shown.'; SP['A2'].font = SUB
hdr(SP, 4, ('№', 'Позиція', 'Item', 'Од. / Unit', 'К-сть / Qty', 'Рядок кошторису / Estimate line', 'Як пораховано / Basis'), 'ABCDEFG'); SP.row_dimensions[4].height = 28
rr = 5; gi = {}
for gk, (gu, ge) in BOM['groups'].items():
    items = [x for x in BOM['items'] if x['g'] == gk]
    if not items: continue
    SP[f'A{rr}'] = f'{gk}. {gu}'; SP[f'C{rr}'] = ge; SP[f'A{rr}'].font = B; SP[f'C{rr}'].font = Font(name=AR, size=9, bold=True, color='4D5C66')
    for c in 'ABCDEFG': SP[f'{c}{rr}'].fill = SEC; SP[f'{c}{rr}'].border = BOX
    rr += 1
    for j, x in enumerate(items, 1):
        SP[f'A{rr}'] = f'{gk}{j}'; SP[f'B{rr}'] = x['uk']; SP[f'C{rr}'] = x['en']; SP[f'D{rr}'] = x['u']; SP[f'E{rr}'] = x['q']; SP[f'F{rr}'] = x['ref']; SP[f'G{rr}'] = x['src']
        SP[f'B{rr}'].font = BLK; SP[f'C{rr}'].font = EN; SP[f'E{rr}'].font = BLUE; SP[f'E{rr}'].number_format = QTY; SP[f'G{rr}'].font = SUB
        for c in 'ABCDEFG': SP[f'{c}{rr}'].border = BOX
        for c in 'BCG': SP[f'{c}{rr}'].alignment = WR
        SP[f'F{rr}'].alignment = Alignment(horizontal='center', vertical='top')
        rr += 1
for c, w in zip('ABCDEFG', (6, 58, 52, 12, 10, 12, 52)): SP.column_dimensions[c].width = w
SP.freeze_panes = 'D5'
# ======================= Стрінги / Strings
T = wb.create_sheet('Стрінги', 2)
T['A1'] = 'Стрінги і MPPT інвертора Deye SUN-100K-G03 / Strings and MPPT of the Deye SUN-100K-G03 inverter'; T['A1'].font = H1
T['A2'] = 'Модуль: Voc 49,5 В, Vmp 41,7 В, Imp 17,51 А, Isc 18,51 А, Voc −0,25 %/K. Інвертор: DC до 1 000 В, старт 250 В, MPPT 200–850 В, 40 А (Isc 60 А) на MPPT. / Module and inverter limits.'; T['A2'].font = SUB
TP = [('Voc модуля, В / Module Voc, V', 49.5), ('Коеф. Voc, %/K / Voc coefficient, %/K', -0.25), ('Мінімальна температура модуля, °C / Min. module temp., °C', -30), ('Vmp при 70 °C, В / Vmp at 70 °C, V', 37.0),
      ('Vmp при −10 °C, В / Vmp at −10 °C, V', 45.93), ('Imp, А', 17.51), ('Isc, А', 18.51), ('Інвертор: макс. DC, В / Inverter max DC, V', 1000), ('Інвертор: MPPT від, В / MPPT min, V', 200),
      ('Інвертор: MPPT до, В / MPPT max, V', 850), ('Інвертор: старт, В / Start-up, V', 250), ('Інвертор: струм на MPPT, А / Current per MPPT, A', 40), ('Інвертор: Isc на MPPT, А / Isc per MPPT, A', 60)]
for i, (nm, v) in enumerate(TP):
    T[f'A{4+i}'] = nm; T[f'B{4+i}'] = v; T[f'A{4+i}'].font = BLK; T[f'B{4+i}'].font = BLUE
    for c in 'AB': T[f'{c}{4+i}'].border = BOX
TR = {i: 4 + i for i in range(len(TP))}; t = lambda i: f'$B${TR[i]}'
r2 = 4 + len(TP) + 1
H2 = ('Перевірка / Check', 'Модулів у стрінгу / Modules', 'Voc −30 °C, В', '≤ 1 000 В', 'Vmp 70 °C, В', '≥ 250 В', 'Vmp −10 °C, В', '≤ 850 В', 'Imp на MPPT, А', '≤ 40 А', 'Isc на MPPT, А', '≤ 60 А')
for c, h in zip('ABCDEFGHIJKL', H2): T[f'{c}{r2}'] = h; T[f'{c}{r2}'].font = B; T[f'{c}{r2}'].fill = HDR; T[f'{c}{r2}'].border = BOX; T[f'{c}{r2}'].alignment = CEN
T.row_dimensions[r2].height = 30
rr = r2 + 1
T[f'A{rr}'] = '«ДБН» / DBN'; T[f'B{rr}'] = f'=Параметри!B{ROW["LS"]}'
T[f'C{rr}'] = f'=B{rr}*{t(0)}*(1+{t(1)}/100*({t(2)}-25))'; T[f'D{rr}'] = f'=IF(C{rr}<={t(7)},"так / yes","ні / no")'
T[f'E{rr}'] = f'=B{rr}*{t(3)}'; T[f'F{rr}'] = f'=IF(AND(E{rr}>={t(8)},E{rr}>={t(10)}),"так / yes","ні / no")'
T[f'G{rr}'] = f'=B{rr}*{t(4)}'; T[f'H{rr}'] = f'=IF(G{rr}<={t(9)},"так / yes","ні / no")'
T[f'I{rr}'] = f'=2*{t(5)}'; T[f'J{rr}'] = f'=IF(I{rr}<={t(11)},"так / yes","ні / no")'
T[f'K{rr}'] = f'=2*{t(6)}'; T[f'L{rr}'] = f'=IF(K{rr}<={t(12)},"так / yes","ні / no")'
for c in 'ABCDEFGHIJKL': T[f'{c}{rr}'].border = BOX; T[f'{c}{rr}'].font = GRN if c == 'B' else BLK
for c in 'CEGIK': T[f'{c}{rr}'].number_format = '#,##0.0'
r3 = rr + 2
PN = {'A': 'арка / arch', 'L': 'нижній дах / lower roof', 'N': 'надбудова / stair house', 'W': 'стіни / walls'}; FN = {'E': 'схід / E', 'W': 'захід / W', 'S': 'південь / S'}
H3 = ('Стрінг / String', 'MPPT', 'Модулів / Modules', 'Де / Where', 'Грані / Faces', 'Нахил, ° / Tilt', 'МВт·год / MWh', 'Невідповідність, % / Mismatch', 'DC «+», м', 'DC «−», м', 'Перемички, м / Jumpers', 'Не працює, модуле-год / Bypassed, module-h')
for c, h in zip('ABCDEFGHIJKL', H3): T[f'{c}{r3}'] = h; T[f'{c}{r3}'].font = B; T[f'{c}{r3}'].fill = HDR; T[f'{c}{r3}'].border = BOX; T[f'{c}{r3}'].alignment = CEN
T.row_dimensions[r3].height = 30
rr = r3 + 1; s0 = rr
off_by = {s['k']: sum(m['offh'] for m in DB['mods'] if m['k'] == s['k']) for s in DB['strings']}
for s in sorted(DB['strings'], key=lambda s: (s['mppt'], s['k'])):
    vals = (f"S{s['k']}", s['mppt'], s['n'], ' + '.join(PN[p] for p in s['parts']), ', '.join(FN[f] for f in s['faces']), s['tilt'], s['mwh'], s['mism'], s['plus']['len'], s['minus']['len'],
            round(sum(j['len'] for j in s['jump']), 1), off_by[s['k']])
    for c, v in zip('ABCDEFGHIJKL', vals): T[f'{c}{rr}'] = v; T[f'{c}{rr}'].border = BOX; T[f'{c}{rr}'].font = BLUE if c in 'CFGHIJKL' else BLK
    T[f'G{rr}'].number_format = '0.000'; T[f'H{rr}'].number_format = '0.0'; T[f'L{rr}'].number_format = '#,##0'
    rr += 1
T[f'A{rr}'] = 'Разом / Total'
for c in 'CGIJKL': T[f'{c}{rr}'] = f'=SUM({c}{s0}:{c}{rr-1})'
for c in 'ABCDEFGHIJKL': T[f'{c}{rr}'].border = BOX; T[f'{c}{rr}'].font = B; T[f'{c}{rr}'].fill = SEC
T[f'G{rr}'].number_format = '0.0'; T[f'L{rr}'].number_format = '#,##0'
rr += 2
T[f'A{rr}'] = ('Не працює — модуль обходиться байпас-діодом, бо його струм нижчий за робочий струм стрінга (часткова тінь або інша орієнтація), за ясного неба, години на рік; '
               'стрінг працює, поки в роботі ≥ 6 модулів (MPPT від 200 В). / Bypassed — the module current is below the string operating current (partial shade or different orientation), clear sky, hours per year.')
T[f'A{rr}'].font = SUB; T.merge_cells(f'A{rr}:L{rr}'); T[f'A{rr}'].alignment = WR; T.row_dimensions[rr].height = 36
for c, w in zip('ABCDEFGHIJKL', (30, 8, 10, 30, 18, 10, 10, 13, 10, 10, 12, 16)): T.column_dimensions[c].width = w
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
os.makedirs(OUT, exist_ok=True)
wb.save(OUT + XN + '.xlsx')
json.dump(dict(KR=KR, ROW=ROW, r2=r2, r3=r3, XN=XN), open(OUT + 'rows_dbn.json', 'w'), ensure_ascii=False)
print('ok', OUT + XN + '.xlsx')
