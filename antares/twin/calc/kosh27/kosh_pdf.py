# -*- coding: utf-8 -*-
"""PDF кошторису 27.10 (дах + ліхтарі + склад) із перерахованого xlsx."""
import openpyxl, asyncio, html
from playwright.async_api import async_playwright
D = '/home/claude/work/kosh27/'
XN = 'Кошторис_будівництва_під_ключ_СЕС_і_УЗЕ_Антарес_дах_і_ліхтарі_27.10.2026'
wb = openpyxl.load_workbook(D + XN + '.xlsx', data_only=True)
K, L, S = wb['Кошторис'], wb['Ліхтарі'], wb['Специфікація — зміни']
e = lambda x: html.escape(str(x)) if x is not None else ''
def nf(x, d=0):
    if x is None or x == '': return ''
    if isinstance(x, str): return e(x)
    s = f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')
    return s
def qty(x):
    if isinstance(x, (int, float)) and abs(x - round(x)) > 1e-9: return nf(x, 2)
    return nf(x)
def pct(x):
    if not isinstance(x, (int, float)): return ''
    v = x * 100; return (nf(v) if abs(v - round(v)) < 1e-9 else nf(v, 1)) + ' %'

# ---- рядки кошторису
rows = []
for r in range(5, 71):
    v = [K.cell(r, c).value for c in range(1, 10)]
    if all(a is None for a in v): continue
    no, nm, u, q, pr, mt, sm, old, bs = v
    if no and nm is None and isinstance(no, str) and no[0].isdigit() and '.' in no[:2]:           # заголовок розділу
        rows.append(f'<tr class="sec"><td colspan="8">{e(no)}</td></tr>'); continue
    if no is None:                                                                                 # підсумки
        cls = 'tot' if nm.startswith('Разом') and 'ПІД КЛЮЧ' not in nm else 'grand'
        if nm.startswith('СОБІВАРТІСТЬ') or 'ПІД КЛЮЧ' in nm: cls = 'grand'
        if nm.startswith('ПДВ') : cls = 'vat'
        if nm.startswith('На 1 кВтп'): rows.append(f'<tr class="tot"><td></td><td colspan="4">{e(nm)}</td><td class="n">{nf(sm)}</td><td class="n o">{nf(old)}</td><td></td></tr>'); continue
        rows.append(f'<tr class="{cls}"><td></td><td colspan="4">{e(nm)}</td><td class="n">{nf(sm)}</td><td class="n o">{nf(old)}</td><td></td></tr>'); continue
    isp = isinstance(q, float) and q < 1 and u is None
    qs = pct(q) if isp else qty(q)
    prs = nf(pr, 2) if isinstance(pr, float) and abs(pr - round(pr)) > 1e-9 and pr < 1000 else nf(pr)
    chg = ''
    if isinstance(sm, (int, float)) and isinstance(old, (int, float)) and abs(sm - old) > 0.5: chg = ' chg'
    if old == 0 and sm == 0 and no == '5.10': chg = ' gap'
    rows.append(f'<tr class="ln{chg}"><td>{e(no)}</td><td>{e(nm)}</td><td class="c">{e(u)}</td><td class="n">{qs}</td><td class="n">{prs}</td>'
                f'<td class="n b">{nf(sm)}</td><td class="n o">{nf(old) if old not in (None, 0) else ("—" if old == 0 else "")}</td><td class="bs">{e(bs)}</td></tr>')
notes = [K.cell(r, 2).value for r in (72, 73, 74)]
g = lambda r: K.cell(r, 7).value; go = lambda r: K.cell(r, 8).value
SUM = [('Закупівля в Китаї', g(16), go(16)), ('Комісія', g(20), go(20)), ('Доставка і мито', g(29), go(29)), ('Купуємо в Україні', g(41), go(41)),
       ('Роботи, фундаменти, EMS', g(54), go(54)), ('Запас і винагорода', g(61), go(61))]
srow = ''.join(f'<tr><td>{a}</td><td class="n">{nf(b)}</td><td class="n o">{nf(c)}</td><td class="n d">{("+" if b - c >= 0 else "−") + nf(abs(b - c))}</td></tr>' for a, b, c in SUM)
# ---- ліхтарі
lr = ''
for r in range(5, 9):
    v = [L.cell(r, c).value for c in range(1, 10)]
    lr += f'<tr><td>{e(v[0])}</td><td>{e(v[1])}</td><td class="c">{e(v[2])}</td><td class="n">{qty(v[3])}</td><td class="n">{pct(v[4])}</td><td class="n b">{qty(v[5])}</td><td class="n">{nf(v[6], 2)}</td><td class="n b">{nf(v[7], 1)}</td><td class="bs">{e(v[8])}</td></tr>'
lr += f'<tr class="tot"><td></td><td colspan="6">Разом кріплення на ліхтарях</td><td class="n">{nf(L["H11"].value)}</td><td></td></tr>'
# ---- специфікація
sr = ''
for r in range(5, 20):
    v = [S.cell(r, c).value for c in range(1, 7)]
    if v[0] is None: continue
    new = v[0].startswith('НОВЕ')
    sr += f'<tr class="{"nw" if new else ""}"><td>{e(v[0].replace("НОВЕ: ", "")[:1].upper() + v[0].replace("НОВЕ: ", "")[1:])}{" <span class=tag>нове</span>" if new else ""}</td><td class="c">{e(v[1])}</td><td class="n o">{qty(v[2]) if v[2] not in (0, None) else "—"}</td><td class="n b">{qty(v[3])}</td><td class="bs">{e(v[4])}</td><td class="bs">{e(v[5])}</td></tr>'

kwp = K['D44'].value; tot = g(63)
H = f'''<!doctype html><html lang="uk"><head><meta charset="utf-8"><title>Кошторис Антарес 27.10</title><style>
@page {{ size:A4 landscape; margin:10mm 11mm 11mm 12mm; }}
body {{ font-family:'DejaVu Sans',sans-serif; font-size:7.4pt; line-height:1.32; color:#1d2a33; margin:0; }}
h1 {{ font-size:12.5pt; margin:0 0 2px; color:#123044; }} h2 {{ font-size:9.6pt; margin:10px 0 4px; color:#123044; }}
.sub {{ color:#5b6b76; font-size:7.3pt; margin-bottom:6px; }}
table {{ border-collapse:collapse; width:100%; }} td, th {{ border-bottom:0.4pt solid #d5dde2; padding:2.2px 4px; vertical-align:top; }}
th {{ background:#e8eef2; text-align:left; font-weight:bold; font-size:7pt; border-bottom:0.8pt solid #9fb1bd; }}
.n {{ text-align:right; white-space:nowrap; }} .c {{ text-align:center; }} .b {{ font-weight:bold; }} .o {{ color:#7a8790; }} .d {{ color:#1f5f78; }}
.bs {{ color:#4d5c66; font-size:6.7pt; }}
tr.sec td {{ background:#f2f5f7; font-weight:bold; color:#123044; padding-top:4px; border-bottom:0.6pt solid #9fb1bd; }}
tr.tot td {{ font-weight:bold; border-bottom:0.8pt solid #9fb1bd; }}
tr.grand td {{ font-weight:bold; background:#e3edf3; font-size:8pt; border-bottom:0.8pt solid #6f8796; }}
tr.vat td {{ color:#4d5c66; }}
tr.chg td.b {{ color:#0b5d86; }} tr.gap td {{ background:#fff6d6; }}
.kpi {{ display:flex; gap:8px; margin:6px 0 8px; }} .k {{ flex:1; border:0.6pt solid #c9d4db; border-radius:4px; padding:5px 7px; }}
.k b {{ display:block; font-size:11pt; color:#123044; }} .k span {{ color:#5b6b76; font-size:6.8pt; }}
.two {{ display:flex; gap:14px; }} .two > div {{ flex:1; }}
.note {{ color:#4d5c66; font-size:6.9pt; margin:3px 0; }}
.tag {{ background:#dff0e6; color:#1d6b40; font-size:6pt; padding:0 3px; border-radius:2px; }}
tr.nw td {{ background:#f6fbf8; }}
.pb {{ page-break-before:always; }}
</style></head><body>
<h1>Кошторис будівництва під ключ — СЕС і УЗЕ 3 × Golen LES-H20, ТОВ МК «Антарес»</h1>
<div class="sub">Дах цеху + модулі на місці ліхтарів + склад. Перерахунок 27.10.2026, $ без ПДВ. Основа — кошторис 26.10 «варіант 1» (1 925 модулів), колонка «Було 26.10» — для порівняння. E-Systems.</div>
<div class="kpi">
<div class="k"><b>2 112 модулів</b><span>дах 1 639 + ліхтарі 297 + склад 176 (було 1 925)</span></div>
<div class="k"><b>{nf(kwp, 2)} кВтп</b><span>192 стрінги, щитів PV 7/1 — 30 + 1 запасний</span></div>
<div class="k"><b>{nf(tot)} $</b><span>під ключ без ПДВ (було {nf(go(63))} $)</span></div>
<div class="k"><b>{nf(g(65))} $</b><span>з ПДВ 20 % (з пільгою «245» — {nf(g(68))} $)</span></div>
<div class="k"><b>{nf(g(70))} $/кВтп</b><span>без ПДВ (було {nf(go(70))} $/кВтп)</span></div>
</div>
<div class="two"><div>
<h2>Підсумок за розділами, $</h2>
<table><tr><th>Розділ</th><th class="n">27.10</th><th class="n">Було 26.10</th><th class="n">Різниця</th></tr>{srow}
<tr class="grand"><td>Собівартість (1–5)</td><td class="n">{nf(g(56))}</td><td class="n o">{nf(go(56))}</td><td class="n d">+{nf(g(56) - go(56))}</td></tr>
<tr class="grand"><td>Разом під ключ без ПДВ</td><td class="n">{nf(tot)}</td><td class="n o">{nf(go(63))}</td><td class="n d">+{nf(tot - go(63))}</td></tr></table>
</div><div>
<h2>Що змінилося проти 26.10</h2>
<div class="note">• Модулі: +297 на місці ліхтарів (вертикально, каркас ≈13°), дах цеху 1 749 → 1 639 (варіант А, v3), склад 176 без змін; до замовлення 2 155 шт. (+2 %).</div>
<div class="note">• Кріплення на ліхтарях — окремою позицією 1.3.2 (рейки 830 м, притиски 629, кронштейни 560); кріплення даху і склад — пропорційно модулям.</div>
<div class="note">• Стрінгів 175 → 192, щитів PV 26 → 30 робочих (по 1 на ліхтар) + 1 запасний; кабелі DC, лотки, лінії Al — пропорційно.</div>
<div class="note">• Баластних блоків менше (6 583 → 6 227): на ліхтарях модулі кріпляться до каркаса без баласту.</div>
<div class="note">• Роботи — за потужністю 50 $/кВт (+6 825 $), механізація — пропорційно модулям.</div>
<div class="note">• Поз. 5.10 «розбирання скління ліхтарів і нова покрівля прорізу» (≈674 м²) — у кошторисі 0 до отримання ціни.</div>
</div></div>
<h2>Кошторис по позиціях</h2>
<table><tr><th style="width:4%">№</th><th style="width:31%">Позиція</th><th class="c" style="width:4%">Од.</th><th class="n" style="width:6%">К-сть</th><th class="n" style="width:7%">Ціна, $</th><th class="n" style="width:8%">Сума, $</th><th class="n" style="width:8%">Було 26.10</th><th>Підстава</th></tr>
{''.join(rows)}</table>
{''.join(f'<div class="note">{e(n)}</div>' for n in notes)}
<h2 style="margin-top:14px">Кріплення модулів на місці ліхтарів (поз. 1.3.2)</h2>
<div class="sub">Вертикально, 1 ряд модулів на схил, 2 рейки вздовж ліхтаря на схил — у зоні притисків ZNSHINE (540 мм від торців модуля). Ціни — КП 26.10 (профіль 41×41 — 1,35 $/м, притиск — 0,50 $), кронштейн — 2 $, крок 1,5 м до замірів каркаса.</div>
<table><tr><th>№</th><th>Позиція</th><th class="c">Од.</th><th class="n">За розкладкою</th><th class="n">Запас</th><th class="n">До замовлення</th><th class="n">Ціна, $</th><th class="n">Сума, $</th><th>Підстава</th></tr>{lr}</table>
<h2>Специфікація замовлення в Китаї — що змінюється проти 26.10</h2>
<div class="sub">Решта позицій специфікації 26.10 (117 позицій, 1 925 модулів) — без змін. Пропорційні кількості — до оновлення КМ і ЕП під розкладку 27.10.</div>
<table><tr><th style="width:30%">Позиція</th><th class="c">Од.</th><th class="n">Було 26.10</th><th class="n">Стає</th><th>Як пораховано</th><th>Уточнити</th></tr>{sr}</table>
</body></html>'''
open(D + 'kosh27.html', 'w', encoding='utf-8').write(H)
async def run():
    async with async_playwright() as p:
        b = await p.chromium.launch(); pg = await b.new_page(); await pg.goto('file://' + D + 'kosh27.html'); await pg.wait_for_timeout(400)
        await pg.pdf(path=D + XN + '.pdf', print_background=True, prefer_css_page_size=True); await b.close()
asyncio.run(run())
print('ok', D + XN + '.pdf')
