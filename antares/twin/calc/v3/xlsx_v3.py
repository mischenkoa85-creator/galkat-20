# -*- coding: utf-8 -*-
"""Таблиця розкладки v3 (xlsx, масштаб 100 %): варіанти, місця модулів, дефлектори, карта опромінення 1 м, джерела."""
import json, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
V = '/home/claude/work/sun/v3/'
J = json.load(open(V + 'roof_v3.json')); L3 = json.load(open(V + 'layout_v3.json')); IR = json.load(open(V + 'irr_v3.json'))
ND = json.load(open(V + 'layout_v3_nodefl.json'))['VARS']; WS_ = json.load(open(V + 'west_v3.json')); DF = json.load(open(V + 'defl_v3.json'))
cand = J['cand']; free = [c for c in cand if c['why'] is None]; key = lambda c: (round(c['xa'], 3), round(c['y'], 3)); fidx = {key(c): i for i, c in enumerate(free)}
mem = []
for V_ in L3['VARS']:
    d = {}
    for s in V_['strings']:
        for i in s['ids']: d[i] = s['id'] + 1
    mem.append(d)
AR = 'Arial'; H = Font(name=AR, bold=True, size=10, color='FFFFFF'); N = Font(name=AR, size=9); B = Font(name=AR, size=9, bold=True)
HF = PatternFill('solid', fgColor='2F4858'); Z = PatternFill('solid', fgColor='F2F5F7'); Y = PatternFill('solid', fgColor='FFF4CC')
th = Side(style='thin', color='C9D1D6'); BD = Border(left=th, right=th, top=th, bottom=th); WR = Alignment(wrap_text=True, vertical='top'); CT = Alignment(horizontal='center', vertical='center', wrap_text=True)
wb = Workbook()
def sheet(ws, head, widths, data, title, note=None, num=None):
    ws.cell(1, 1, title).font = Font(name=AR, bold=True, size=12); r0 = 2
    if note:
        ws.cell(2, 1, note).font = Font(name=AR, size=9, italic=True, color='5B6770'); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(head)); ws.cell(2, 1).alignment = WR; ws.row_dimensions[2].height = 48; r0 = 3
    for j, h in enumerate(head, 1):
        c = ws.cell(r0, j, h); c.font = H; c.fill = HF; c.alignment = CT; c.border = BD
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    for i, row in enumerate(data):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + 1 + i, j, v); c.font = N; c.alignment = WR; c.border = BD
            if num and j in num and isinstance(v, (int, float)): c.number_format = num[j]
            if i % 2: c.fill = Z
    ws.freeze_panes = ws.cell(r0 + 1, 1); ws.sheet_view.zoomScale = 100
    if data: ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(head))}{r0 + len(data)}'
    return r0
# 1. Варіанти
ws = wb.active; ws.title = 'Варіанти'
VS = L3['VARS']; LV = L3['LAN']
rows = [[f"{nm}", v['sel'], v['n'], v['st'], v['kwp'], v['mwh'], v['y'], v['irr'], v['mism'], v['cost'], v['rev'], v['pb'], note]
        for nm, v, note in ((f'А — усі вільні місця (рекомендовано)', VS[0], 'Найбільше енергії; кожен шар місць окупається'),
                            ('Б — місця від 1 050 кВт·год/м²', VS[1], 'Без смуги біля західної будівлі й найбільш затінених місць'),
                            ('В — місця від 1 100 кВт·год/м²', VS[2], 'Лише найкращі м²')) ]
rows.append(['А без смуги 7,5 м біля західної будівлі', '', WS_['n'], WS_['st'], WS_['kwp'], WS_['mwh'], WS_['y'], '', '', round(WS_['n'] * L3['COST']), '', '', f"Смуга: {WS_['dn']} модулів, {WS_['de']} МВт·год, {WS_['yd']} кВт·год/кВтп, окупність ≈{WS_['pb']} р.; снігова зона біля вищої стіни"])
rows.append(['А, якщо дефлектори приберуть (як у розрахунку 26.10)', ND[0]['sel'], ND[0]['n'], ND[0]['st'], ND[0]['kwp'], ND[0]['mwh'], ND[0]['y'], ND[0]['irr'], ND[0]['mism'], ND[0]['cost'], ND[0]['rev'], ND[0]['pb'], 'Для порівняння: що забирають дефлектори'])
rows.append(['+ модулі книжково на ліхтарях 3,66 × 1,89 (окремо)', LV['places'], LV['n'], LV['st'], LV['kwp'], LV['mwh'], LV['y'], '', '', round(LV['n'] * L3['COST']), '', LV['pb'], 'Ліхтарі світлопрозорі — цех стане темнішим; рішення власника'])
r0 = sheet(ws, ['Варіант', 'Вільних місць', 'Модулів у стрінгах', 'Стрінгів × 11', 'кВтп', 'МВт·год за рік', 'кВт·год/кВтп', 'Опромінення модулів, кВт·год/м²', 'Різна освітленість, %', 'Вартість модулів під ключ, $ (оцінка)', 'Ефект за рік, $', 'Окупність цієї частини, р.', 'Примітка'],
           [44, 10, 11, 9, 10, 11, 10, 13, 11, 14, 11, 11, 48], rows, 'Дах цеху Антарес — варіанти розкладки за картою опромінення (v3, 27.10.2026)',
           'Модулі ZNSHINE 730 Вт на гранях покрівлі (розріз 1-1), стрінги по 11 у межах однієї грані. Генерація — модель двійника (Meteonorm 8.1, крок 10 хв, тіні від усього на даху), на виході PV до PCS. '
           'Ціна 1 кВт·год 7,00 грн з ПДВ (оцінка власника), ККД PCS 97,5 %, 44,89 грн/$; ≈250 $ на модуль під ключ без ПДВ — частина СЕС кошторису 26.10 без УЗЕ, EMS, фундаментів і мідних ліній AC (оцінка).',
           num={5: '#,##0.0', 6: '#,##0.0', 10: '#,##0', 11: '#,##0', 12: '0.0'})
for j in range(1, 14): ws.cell(r0 + 1, j).fill = Y
# 2. Місця
ws2 = wb.create_sheet('Місця модулів')
WHYT = {'lan': 'ліхтар + 0,25 м', 'defl': 'дефлектор / труба + 0,3 м', 'ac': 'рама кондиціонерів + 1 м', 'exit': 'вихід, портал, прохід до рами', 'door': 'прохід від дверей західної будівлі', 'drain': 'воронка'}
data = []
for n, c in enumerate(cand, 1):
    i = fidx.get(key(c))
    data.append([n, c['sp'], 'захід' if c['face'] == 'W' else 'схід', c['tilt'], round(c['xa'], 2), round(c['xb'], 2), round(c['y'], 3), WHYT.get(c['why'], '') if c['why'] else '',
                 c.get('kwh') if i is not None else None, c.get('loss') if i is not None else None, c.get('irr') if i is not None else None,
                 mem[0].get(i) if i is not None else None, mem[1].get(i) if i is not None else None, mem[2].get(i) if i is not None else None])
r0 = sheet(ws2, ['№', 'Проліт (осі)', 'Схил на', 'Кут, °', 'X від, м', 'X до, м', 'Y центр, м', 'Місце зайняте', 'кВт·год за рік', 'Тінь, %', 'Опромінення, кВт·год/м²', 'Стрінг у А', 'Стрінг у Б', 'Стрінг у В'],
           [6, 9, 8, 7, 9, 9, 9, 26, 10, 8, 12, 9, 9, 9], data, 'Усі місця під модулі на гранях покрівлі',
           'X — на схід від західного краю даху, Y — на північ від південного краю (як у 3D-двійнику). Модуль 2 384 × 1 303 мм довгою стороною вздовж Y. Порожній стрінг — місце не входить у варіант.',
           num={5: '0.00', 6: '0.00', 7: '0.000', 9: '#,##0.0', 10: '0.0', 11: '#,##0'})
ws2.conditional_formatting.add(f'K{r0 + 1}:K{r0 + len(data)}', ColorScaleRule(start_type='num', start_value=850, start_color='B4461B', mid_type='num', mid_value=1050, mid_color='F6C453', end_type='num', end_value=1185, end_color='FFF3B0'))
# 3. Дефлектори
ws3 = wb.create_sheet('Дефлектори')
data = [[o['n'], o['x'], o['y'], o['d'], o['h'], o['src'], o['st']] for o in DF]
sheet(ws3, ['№', 'X, м', 'Y, м', 'Діаметр, м', 'Висота, м', 'Джерело', 'Стан у розрахунку'], [6, 9, 9, 10, 10, 34, 44], data, 'Дефлектори, «грибочки» й труби на даху цеху',
      'Положення — із супутникового знімка Google (знімок до 2026, точність ≈0,5 м) і РП. Розмір і висота — з РП або типові 1,2 × 1,6 м, де на знімку видно лише сам об\'єкт. «Уточнити» — малі труби з РП, на знімку не видно. Номери — як на плані в PDF і в описі дефлектора в 3D.',
      num={2: '0.0', 3: '0.0', 4: '0.00', 5: '0.0'})
# 4. Карта 1 м
ws4 = wb.create_sheet('Карта 1 м')
nx, ny, res = IR['nx'], IR['ny'], IR['res']; E = np.array([np.nan if (v is None or k != 'r') else v for v, k in zip(IR['e'], IR['k'])], float).reshape(ny, nx)
m1 = np.full((ny // 2, nx // 2), np.nan)
for j in range(ny // 2):
    for i in range(nx // 2):
        b = E[2 * j:2 * j + 2, 2 * i:2 * i + 2]; m1[j, i] = np.nanmean(b) if np.isfinite(b).any() else np.nan
ws4.cell(1, 1, 'Опромінення покрівлі, кВт·год/м² за рік, клітинки 1 × 1 м (середнє з 4 клітинок 0,5 м); рядки — Y від півночі до півдня, стовпці — X із заходу на схід. Порожньо — ліхтар, рама кондиціонерів, дефлектор.').font = Font(name=AR, size=9, italic=True)
ws4.cell(3, 1, 'Y \\ X').font = B
for i in range(m1.shape[1]): c = ws4.cell(3, i + 2, i + 0.5); c.font = Font(name=AR, size=7, bold=True); ws4.column_dimensions[get_column_letter(i + 2)].width = 4.2
ws4.column_dimensions['A'].width = 6
for r, j in enumerate(range(m1.shape[0] - 1, -1, -1)):
    ws4.cell(4 + r, 1, j + 0.5).font = Font(name=AR, size=7, bold=True)
    for i in range(m1.shape[1]):
        v = m1[j, i]
        if np.isfinite(v): c = ws4.cell(4 + r, i + 2, int(round(v))); c.font = Font(name=AR, size=6)
ws4.conditional_formatting.add(f'B4:{get_column_letter(m1.shape[1] + 1)}{3 + m1.shape[0]}', ColorScaleRule(start_type='num', start_value=800, start_color='6B2A6E', mid_type='num', mid_value=1050, mid_color='E8892A', end_type='num', end_value=1185, end_color='FFF3B0'))
ws4.freeze_panes = 'B4'; ws4.sheet_view.zoomScale = 100
# 5. Джерела
ws5 = wb.create_sheet('Джерела і заміри')
data = [['Профіль покрівлі: грані по ≈3 м, ендови +15,90, гребені +17,66 / +18,15', 'Розріз 1-1 (DWG замовника, 26.10)', 'висока'],
        ['Ліхтарі: центри по X, 4 шт. на прольотах 18 м', 'Розріз 1-1', 'висока'], ['Ліхтарі: довжина 23,7–78,4 м (ліхтар 1 — 30,6–78,4)', 'РП; збігається із супутником ±0,5 м', 'висока'],
        ['Ліхтарі: гребінь ≈1,89 м над покрівлею, по низу ≈3,66 м, схили світлопрозорі', 'Рулетка й фото 27.10', 'висока'],
        ['Торцеві парапети 1,37 м (південний — в ендові; на уступах ≈1,65)', 'Рулетка 27.10; північний не міряли — прийнято як південний', 'середня'],
        ['Дефлектори і труби: 49 + 8 малих', 'Супутник Google + РП', 'середня — знімок до 2026'],
        ['Рама кондиціонерів 14,5 × 2,4 × 2,2 м', 'Власник 27.10; положення — фото, ±2 м', 'низька (положення)'],
        ['Вихід на дах і портал', 'Фото 26–27.10 і супутник: X ≈62–70 м', 'низька'], ['Воронка біля південного парапету, центр ≈0,87 м від стіни', 'Рулетка 27.10', 'середня'],
        ['Західна будівля ≈6,8 м над дахом', 'Фото (кладка)', 'середня; ширина невідома'], ['Поворот будівлі ≈3,8° за годинниковою стрілкою', 'Супутник; у моделі 4°', 'висока'],
        ['Клімат', 'Meteonorm 8.1 з PVsyst (як у РП)', 'висока']]
sheet(ws5, ['Що', 'Джерело', 'Надійність'], [70, 50, 26], data, 'Джерела і заміри, на яких зроблено v3',
      'Перевірити рулеткою: рама кондиціонерів від виходу; положення виходу й порталу від східного краю; північний парапет; які дефлектори лишаються після заміни покрівлі.')
OUT = V + 'Таблиця_розкладка_даху_за_картою_опромінення_Антарес_27.10.2026.xlsx'
wb.save(OUT); print(OUT)
