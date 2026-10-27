# -*- coding: utf-8 -*-
"""Будує Відомість фото (xlsx) і список переміщень на Диску. → out/"""
import json, os, sys
sys.path.insert(0, os.path.dirname(__file__))
from reg_data import F, OBJ, P27, WAS27, WAS_OVERRIDE, G2F, G2O, OVR26, MEAS
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
SP = '/tmp/claude-0/-home-claude/a46e64fa-3229-5beb-8862-98d878ef89e1/scratchpad/ph/'
SUB = json.load(open(SP + 'sub.json')); ROOT = json.load(open(SP + 'root.json'))
s = open('/home/claude/work/twin/foto.js', encoding='utf-8').read(); FOTO = json.loads(s[s.find('{'):s.rfind('}') + 1])
TWIN27 = set(json.load(open('/home/claude/work/photos27/twin27.json'))) if os.path.exists('/home/claude/work/photos27/twin27.json') else set()
FIDX27 = {p[0]: i for i, p in enumerate(FOTO['p']) if len(p) > 7}
ids27 = {}
for f in SUB: ids27.setdefault(f['t'][9:15], []).append(f['id'])
rows = []   # dict
for t, fk, ob, x, y, az, d, m in P27:
    was = WAS_OVERRIDE.get(t, WAS27[fk if fk != 'yard' else 'abk'])
    if was == 'Певнічний фасад': was = 'Певнічний фасад (перейменовано)'
    ii = ids27.get(t, [])
    rows.append(dict(f=f'20261027_{t}.jpg', date='27.10.2026', t=f'{t[:2]}:{t[2:4]}:{t[4:]}', was=was, now=F[fk][0], fk=fk, ob=ob, x=x, y=y, az=az, d=d + ('' if len(ii) < 2 else ''), m=m,
                     twin=(f'так, № {FIDX27[t] + 1}' if t in FIDX27 else 'ні') if x is not None else '—', id=ii[0] if ii else '', dup=len(ii) > 1))
fidx = {p[0]: i for i, p in enumerate(FOTO['p'])}
for f in sorted(ROOT, key=lambda r: r['t']):
    t = f['t'][9:15]; ext = f['t'].split('.')[-1]
    if t in fidx:
        p = FOTO['p'][fidx[t]]; fk, ob = OVR26.get(t, (G2F[p[4]], G2O[p[4]]))
        d = p[5] + (' (відео 2 с)' if ext == 'mp4' else '')
        rows.append(dict(f=f['t'], date='26.10.2026', t=f'{t[:2]}:{t[2:4]}:{t[4:]}', was=F['root'][0], now=F[fk][0], fk=fk, ob=ob, x=p[1], y=p[2], az=p[3], d=d, m='',
                         twin=f'так, № {fidx[t] + 1}', id=f['id'], dup=False))
    else:   # відео 12:37 на даху
        rows.append(dict(f=f['t'], date='26.10.2026', t=f'{t[:2]}:{t[2:4]}:{t[4:]}', was=F['root'][0], now=F['roof'][0], fk='roof', ob='O01', x=None, y=None, az=None,
                         d='Відео з даху цеху (11,7 МБ), між східним парапетом і ПнЗх кутом — не переглянуто', m='', twin='ні', id=f['id'], dup=False))
rows.sort(key=lambda r: (r['date'][3:5] + r['date'][:2], r['t']))
moves = [r for r in rows if r['was'] != r['now'] and r['id'] and not r['was'].startswith('Певнічний')]
json.dump([dict(f=r['f'], id=r['id'], was=r['was'], now=r['now'], pid=F[r['fk']][1]) for r in moves], open('/home/claude/work/photos27/moves.json', 'w'), ensure_ascii=False, indent=0)
json.dump(rows, open('/home/claude/work/photos27/rows.json', 'w'), ensure_ascii=False)
# ---------------- xlsx
AR = 'Arial'; H = Font(name=AR, bold=True, size=10, color='FFFFFF'); N = Font(name=AR, size=9); B = Font(name=AR, size=9, bold=True); L = Font(name=AR, size=9, color='1F4E9E', underline='single')
HF = PatternFill('solid', fgColor='2F4858'); Z = PatternFill('solid', fgColor='F2F5F7'); Y = PatternFill('solid', fgColor='FFF4CC'); R = PatternFill('solid', fgColor='FBE3DC')
th = Side(style='thin', color='C9D1D6'); BD = Border(left=th, right=th, top=th, bottom=th); WR = Alignment(wrap_text=True, vertical='top'); CT = Alignment(horizontal='center', vertical='top', wrap_text=True)
wb = Workbook()
def sheet(ws, head, widths, data, title=None, note=None):
    r0 = 1
    if title: ws.cell(1, 1, title).font = Font(name=AR, bold=True, size=12); r0 = 2
    if note: ws.cell(r0, 1, note).font = Font(name=AR, size=9, italic=True, color='5B6770'); ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=len(head)); ws.cell(r0, 1).alignment = WR; ws.row_dimensions[r0].height = 42; r0 += 1
    for j, h in enumerate(head, 1):
        c = ws.cell(r0, j, h); c.font = H; c.fill = HF; c.alignment = CT; c.border = BD
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    for i, row in enumerate(data):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + 1 + i, j, v); c.font = N; c.alignment = WR; c.border = BD
            if i % 2: c.fill = Z
    ws.freeze_panes = ws.cell(r0 + 1, 1); ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(head))}{r0 + len(data)}'
    ws.sheet_view.zoomScale = 100
    return r0
ws = wb.active; ws.title = 'Відомість фото'
head = ['№', 'Файл', 'Дата', 'Час', 'Папка була', 'Папка стала', 'Об\'єкт', 'Назва об\'єкта', 'X, м', 'Y, м', 'Напрям, °', 'Що на фото', 'Виміри з фото', 'Об\'єкт у 3D', 'Точка в двійнику', 'Відкрити']
data = []
for i, r in enumerate(rows, 1):
    data.append([i, r['f'], r['date'], r['t'], r['was'], r['now'], r['ob'], OBJ[r['ob']][0], r['x'], r['y'], r['az'], r['d'] + (' · на Диску 2 однакові копії' if r['dup'] else ''), r['m'], OBJ[r['ob']][1], r['twin'], 'Диск' if r['id'] else ''])
r0 = sheet(ws, head, [5, 21, 10, 8, 20, 22, 7, 26, 6, 6, 8, 58, 30, 10, 12, 8], data, 'Відомість фото обстеження — Антарес, 26–27.10.2026',
           'X — на схід від західного краю даху цеху, Y — на північ від південного краю, м (як у 3D-двійнику); напрям — азимут погляду (0 — північ, 90 — схід). Точки 27.10 — оцінка за фото: ±3–10 м, на даху АБК ±10 м (розміри АБК не виміряні). Порожньо — не визначено. Жовтим — є виміри; червоним — файл не завантажився. Фото з номерами авто і договорів — лише приватно (Диск, приватний двійник), не на GitHub.')
for i, r in enumerate(rows):
    rr = r0 + 1 + i
    if r['id']:
        c = ws.cell(rr, 16); c.hyperlink = f'https://drive.google.com/file/d/{r["id"]}/view'; c.font = L
    if r['m']:
        for j in (12, 13): ws.cell(rr, j).fill = Y
    if 'Не завантажилось' in r['d']:
        for j in range(1, 17): ws.cell(rr, j).fill = R
    if r['was'] != r['now'] and not r['was'].startswith('Певнічний'): ws.cell(rr, 6).font = B
# об'єкти
ws2 = wb.create_sheet('Об\'єкти і 3D')
cnt = {}
for r in rows: cnt.setdefault(r['ob'], []).append(r['f'][:15].replace('2026', '').replace('_', ' ').strip())
data = [[k, v[0], v[1], v[2], v[3], len(cnt.get(k, [])), ', '.join(sorted(set(x[-6:] for x in cnt.get(k, []))))] for k, v in OBJ.items()]
r0b = sheet(ws2, ['Код', 'Об\'єкт', 'У 3D', 'Зроблено в двійнику 27.10', 'Ще зробити / уточнити', 'Фото', 'Фото (час)'], [6, 28, 10, 42, 56, 6, 50], data, 'Об\'єкти: що є в 3D-двійнику і чого немає',
            'Порівняно з двійником (приватна версія з фото оновлена 27.10). «частково» / «є умовно» — у моделі неповно або за оцінкою; виміри — аркуш «Виміри».')
for i, row in enumerate(data):
    if row[2] in ('немає', 'частково', 'є умовно'):
        ws2.cell(r0b + 1 + i, 3).fill = Y
# виміри
ws3 = wb.create_sheet('Виміри')
sheet(ws3, ['Що', 'Значення', 'Фото', 'У моделі зараз', 'Примітка'], [40, 22, 24, 26, 50], [list(m) for m in MEAS], 'Виміри, прочитані з фото (рулетка)',
      'Значення прочитано з фото — точність ±1–2 см за шкалою, але де нуль рулетки, видно не на всіх фото. Варто підтвердити рулеткою там, де стоїть «≈».')
# папки
ws4 = wb.create_sheet('Папки')
fold = {}
for r in rows: fold.setdefault(r['now'], [0, 0]); fold[r['now']][0 if r['date'].startswith('26') else 1] += 1
what = {'Дах цеху': 'Покрівля цеху загалом: вихід, портал, рама кондиціонерів, дефлектори, західна стіна, східний парапет', 'Фонарі на даху цеха': 'Ліхтарі: заміри, прозорі схили, вигляд усередину',
        'Борт на даху': 'Південний парапет: висота, уступи, кабель', 'Лівневки на даху': 'Водостічні воронки: положення, ковпаки', 'Дах абк': 'Дах адмінкорпусу: вентшахти, вентвиходи, антени; види на цех і стадіон',
        'Південний фасад і проїзд': 'Фасад зі сходами на дах, проїзд, контейнер, будки, кондиціонери на землі (була «Певнічний фасад»)', 'Північний фасад': 'Північний фасад, ТП-6923, цегляна прибудова, в\'їзд',
        'Східний фасад': 'Східний фасад, ворота, прибудова ТП-5913', 'Двір, ТП-5913 і адмінкорпус': 'Проїзд між ТП-5913 і АБК, камери трансформаторів, АБК з двору', 'ТП усередині': 'Схеми КТП-5913, РУ-0,4 кВ, кабелі в ТП'}
data = [[k, what.get(k, ''), v[0], v[1], v[0] + v[1], 'Диск'] for k, v in sorted(fold.items(), key=lambda kv: list(what).index(kv[0]) if kv[0] in what else 99)]
r0d = sheet(ws4, ['Папка', 'Що в ній', '26.10', '27.10', 'Разом', 'Відкрити'], [28, 70, 7, 7, 7, 9], data, 'Папки в «Фото антаріс» після розкладання',
            'Нові папки: «Північний фасад», «Східний фасад», «Двір, ТП-5913 і адмінкорпус», «ТП усередині». Папку «Певнічний фасад» перейменовано на «Південний фасад і проїзд» (на фото — фасад зі сходами на дах; у моделі це південь).')
byname = {v[0]: v[1] for v in F.values()}
for i, row in enumerate(data):
    c = ws4.cell(r0d + 1 + i, 6); c.hyperlink = f'https://drive.google.com/drive/folders/{byname[row[0]]}'; c.font = L
# переміщення (журнал і інструкція)
ws5 = wb.create_sheet('Переміщення')
data = [[i, r['f'], r['was'], r['now'], 'перенесено'] for i, r in enumerate(moves, 1)]
sheet(ws5, ['№', 'Файл', 'З папки', 'У папку', 'Стан'], [5, 24, 28, 28, 14], data, 'Що куди перенесено на Диску (27.10)',
      'Перенесено інструментом Диска 27.10. Якщо треба повторити вручну або іншим помічником: у «Фото антаріс» створити 4 папки з аркуша «Папки», перейменувати «Певнічний фасад» → «Південний фасад і проїзд», далі кожен файл з колонки «Файл» перенести в «У папку». Файли не перейменовувати.')
os.makedirs('/home/claude/work/photos27/out', exist_ok=True)
OUT = '/home/claude/work/photos27/out/Відомість_фото_обстеження_Антарес_26-27.10.2026.xlsx'
wb.save(OUT); print(len(rows), 'rows;', len(moves), 'moves;', OUT)
